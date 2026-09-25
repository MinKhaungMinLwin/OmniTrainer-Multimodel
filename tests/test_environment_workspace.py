from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from services.api.omni_api.config import Settings
from services.api.omni_api.main import create_app


def build_app(database_path: Path, attachment_dir: Path):
    return create_app(
        Settings(
            environment="test",
            database_url=f"sqlite+aiosqlite:///{database_path}",
            jwt_secret="environment-test-secret-" * 2,
            allow_dev_auth=True,
            auto_create_schema=True,
            storage_backend="local",
            attachment_dir=str(attachment_dir),
        )
    )


def authenticate(client: TestClient) -> dict[str, str]:
    token = client.post("/api/v1/dev/token", json={"email": "owner@omni.example"}).json()["access_token"]
    auth = {"Authorization": f"Bearer {token}"}
    tenant_id = client.get("/api/v1/tenants", headers=auth).json()[0]["id"]
    return {**auth, "X-Tenant-ID": tenant_id, "X-Correlation-ID": "environment-test"}


def laboratory_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["sample_id", "analyte", "result", "unit", "reporting_limit", "qualifier"])
    sheet.append(["BH01", "Lead", 420, "mg/kg", 1, None])
    sheet.append(["BH02", "Benzene", "<0.01", "mg/L", 0.01, "<"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_environmental_workbook_report_review_and_exports(tmp_path: Path):
    with TestClient(build_app(tmp_path / "environment.db", tmp_path / "attachments")) as client:
        headers = authenticate(client)
        project_response = client.post(
            "/api/v1/environment/projects",
            headers=headers,
            json={"name": "Riverside assessment", "code": "riverside", "jurisdiction": "Western Australia"},
        )
        assert project_response.status_code == 201
        project = project_response.json()

        upload = client.post(
            f"/api/v1/environment/projects/{project['id']}/workbooks",
            headers=headers,
            files={
                "file": (
                    "laboratory.xlsx",
                    laboratory_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert upload.status_code == 201
        workbook = upload.json()
        assert workbook["status"] == "validated"
        assert workbook["validation"]["row_count"] == 2
        assert sum(workbook["validation"]["summary"].values()) == 2

        search = client.get(
            "/api/v1/environment/documents/search",
            headers=headers,
            params={"q": "laboratory screening review"},
        )
        assert search.status_code == 200
        assert search.json()["abstained"] is False
        assert search.json()["matches"][0]["source_uri"].startswith("environment://")

        draft_response = client.post(
            f"/api/v1/environment/projects/{project['id']}/reports",
            headers=headers,
            json={
                "workbook_id": workbook["id"],
                "question": "Screen the laboratory results and explain the required review.",
            },
        )
        assert draft_response.status_code == 201
        draft = draft_response.json()
        assert draft["status"] == "draft"
        assert "Portfolio demonstration only" in draft["report_markdown"]

        draft_docx = client.get(f"/api/v1/environment/reports/{draft['id']}/export.docx", headers=headers)
        assert draft_docx.status_code == 200
        document = Document(BytesIO(draft_docx.content))
        assert "DRAFT — NOT APPROVED" in "\n".join(paragraph.text for paragraph in document.paragraphs)

        approved_response = client.post(
            f"/api/v1/environment/reports/{draft['id']}/review",
            headers=headers,
            json={"decision": "approve", "reason": "Checked against the synthetic certificate and criteria."},
        )
        assert approved_response.status_code == 200
        assert approved_response.json()["status"] == "approved"

        approved_xlsx = client.get(f"/api/v1/environment/reports/{draft['id']}/export.xlsx", headers=headers)
        assert approved_xlsx.status_code == 200
        exported = load_workbook(BytesIO(approved_xlsx.content), read_only=True)
        assert exported["Report summary"]["A1"].value == "APPROVED"
        assert exported["Screening results"].max_row == 3

        assert len(client.get("/api/v1/environment/projects", headers=headers).json()) == 1
        assert len(client.get(f"/api/v1/environment/projects/{project['id']}/reports", headers=headers).json()) == 1


def test_invalid_environmental_workbook_cannot_generate_report(tmp_path: Path):
    with TestClient(build_app(tmp_path / "invalid.db", tmp_path / "attachments")) as client:
        headers = authenticate(client)
        project = client.post(
            "/api/v1/environment/projects",
            headers=headers,
            json={"name": "Invalid data", "code": "invalid-data", "jurisdiction": "Thailand"},
        ).json()
        book = Workbook()
        book.active.append(["wrong", "columns"])
        output = BytesIO()
        book.save(output)
        workbook = client.post(
            f"/api/v1/environment/projects/{project['id']}/workbooks",
            headers=headers,
            files={"file": ("invalid.xlsx", output.getvalue())},
        ).json()
        assert workbook["status"] == "invalid"
        response = client.post(
            f"/api/v1/environment/projects/{project['id']}/reports",
            headers=headers,
            json={"workbook_id": workbook["id"], "question": "Create a report"},
        )
        assert response.status_code == 422
