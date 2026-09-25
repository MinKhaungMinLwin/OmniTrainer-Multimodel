import json
import os
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, Field


DISCLAIMER = (
    "Portfolio demonstration only. Screening results require review by a qualified "
    "environmental practitioner before they are used for a regulatory or client decision."
)
REQUIRED_COLUMNS = {"sample_id", "analyte", "result", "unit", "reporting_limit"}


class LabResult(BaseModel):
    row: int
    sample_id: str
    analyte: str
    result: float
    unit: str
    reporting_limit: float
    qualifier: str | None = None


class ScreeningCriterion(BaseModel):
    analyte: str
    value: float
    unit: str
    source_uri: str
    label: str


class ScreenedResult(BaseModel):
    sample_id: str
    analyte: str
    result: float
    unit: str
    criterion: float | None = None
    criterion_unit: str | None = None
    classification: str
    source_uri: str | None = None


class WorkbookValidation(BaseModel):
    project_id: str
    file: str
    valid: bool
    row_count: int
    errors: list[str] = Field(default_factory=list)
    results: list[LabResult] = Field(default_factory=list)
    screening: list[ScreenedResult] = Field(default_factory=list)
    summary: dict[str, int] = Field(default_factory=dict)
    disclaimer: str = DISCLAIMER


def workspace_root() -> Path:
    configured = os.getenv("OMNI_ENVIRONMENT_WORKSPACE", "data/environment")
    return Path(configured).expanduser().resolve()


def _safe_slug(value: str, label: str) -> str:
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,63}", value):
        raise ValueError(f"{label} must contain only lowercase letters, numbers, and hyphens")
    return value


def project_directory(project_id: str, root: Path | None = None) -> Path:
    safe_id = _safe_slug(project_id, "project_id")
    base = (root or workspace_root()).resolve()
    return (base / "projects" / safe_id).resolve()


def project_file(project_id: str, relative_path: str, root: Path | None = None) -> Path:
    project_root = project_directory(project_id, root)
    candidate = (project_root / relative_path).resolve()
    try:
        candidate.relative_to(project_root)
    except ValueError:
        raise ValueError("file must stay inside the selected project workspace") from None
    if candidate.suffix.lower() != ".xlsx":
        raise ValueError("laboratory input must be an .xlsx workbook")
    if not candidate.is_file():
        raise ValueError(f'workbook "{relative_path}" was not found')
    if candidate.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("laboratory workbook exceeds the 20 MB safety limit")
    return candidate


def load_criteria(root: Path | None = None) -> list[ScreeningCriterion]:
    path = (root or workspace_root()) / "standards" / "demo-screening-criteria.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    return [ScreeningCriterion.model_validate(item) for item in payload["criteria"]]


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _number(value: Any, row: int, field: str) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value or "").strip()
    if field == "result":
        text = text.lstrip("<>").strip()
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"row {row}: {field} must be numeric") from None


def _convert(value: float, source_unit: str, target_unit: str) -> float | None:
    if source_unit == target_unit:
        return value
    conversions = {
        ("ug/L", "mg/L"): 0.001,
        ("mg/L", "ug/L"): 1000.0,
        ("ug/kg", "mg/kg"): 0.001,
        ("mg/kg", "ug/kg"): 1000.0,
    }
    factor = conversions.get((source_unit, target_unit))
    return value * factor if factor is not None else None


def _screen(result: LabResult, criteria: list[ScreeningCriterion]) -> ScreenedResult:
    criterion = next((item for item in criteria if item.analyte.casefold() == result.analyte.casefold()), None)
    if criterion is None:
        return ScreenedResult(
            sample_id=result.sample_id,
            analyte=result.analyte,
            result=result.result,
            unit=result.unit,
            classification="no_criterion",
        )
    comparable = _convert(result.result, result.unit, criterion.unit)
    if comparable is None:
        return ScreenedResult(
            sample_id=result.sample_id,
            analyte=result.analyte,
            result=result.result,
            unit=result.unit,
            criterion=criterion.value,
            criterion_unit=criterion.unit,
            classification="unit_review_required",
            source_uri=criterion.source_uri,
        )
    return ScreenedResult(
        sample_id=result.sample_id,
        analyte=result.analyte,
        result=result.result,
        unit=result.unit,
        criterion=criterion.value,
        criterion_unit=criterion.unit,
        classification="above_criterion" if comparable > criterion.value else "at_or_below_criterion",
        source_uri=criterion.source_uri,
    )


def validate_lab_workbook(
    project_id: str,
    relative_path: str,
    *,
    root: Path | None = None,
) -> WorkbookValidation:
    path = project_file(project_id, relative_path, root)
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [_normalise_header(value) for value in next(rows, ())]
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        return WorkbookValidation(
            project_id=project_id,
            file=relative_path,
            valid=False,
            row_count=0,
            errors=[f"missing required columns: {', '.join(missing)}"],
        )
    indexes = {name: headers.index(name) for name in headers if name}
    results: list[LabResult] = []
    errors: list[str] = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        record = {name: values[index] if index < len(values) else None for name, index in indexes.items()}
        try:
            result_text = str(record["result"] or "").strip()
            qualifier = str(record.get("qualifier") or "").strip() or None
            if not qualifier and result_text[:1] in {"<", ">"}:
                qualifier = result_text[:1]
            results.append(
                LabResult(
                    row=row_number,
                    sample_id=str(record["sample_id"] or "").strip(),
                    analyte=str(record["analyte"] or "").strip(),
                    result=_number(record["result"], row_number, "result"),
                    unit=str(record["unit"] or "").strip(),
                    reporting_limit=_number(record["reporting_limit"], row_number, "reporting_limit"),
                    qualifier=qualifier,
                )
            )
        except (ValueError, TypeError) as exc:
            errors.append(str(exc))
    criteria = load_criteria(root)
    screened = [_screen(result, criteria) for result in results]
    classifications: dict[str, int] = {}
    for item in screened:
        classifications[item.classification] = classifications.get(item.classification, 0) + 1
    return WorkbookValidation(
        project_id=project_id,
        file=relative_path,
        valid=not errors and bool(results),
        row_count=len(results),
        errors=errors or ([] if results else ["workbook contains no laboratory result rows"]),
        results=results,
        screening=screened,
        summary=classifications,
    )


def search_documents(query: str, *, root: Path | None = None, limit: int = 5) -> dict[str, Any]:
    terms = {term for term in re.findall(r"[a-z0-9]+", query.casefold()) if len(term) > 2}
    documents = (root or workspace_root()) / "documents"
    matches: list[dict[str, Any]] = []
    for path in sorted(documents.glob("*.md")):
        content = path.read_text(encoding="utf-8")
        score = sum(content.casefold().count(term) for term in terms)
        if score:
            paragraph = next(
                (part.strip() for part in content.split("\n\n") if any(term in part.casefold() for term in terms)),
                content[:500].strip(),
            )
            matches.append(
                {
                    "title": content.splitlines()[0].lstrip("# ").strip(),
                    "source_uri": f"environment://documents/{path.name}",
                    "score": score,
                    "snippet": paragraph[:700],
                }
            )
    matches.sort(key=lambda item: (-item["score"], item["source_uri"]))
    return {"query": query, "matches": matches[:limit], "abstained": not matches}


def draft_screening_report(
    project_id: str,
    lab_file: str,
    question: str,
    *,
    root: Path | None = None,
) -> dict[str, Any]:
    validation = validate_lab_workbook(project_id, lab_file, root=root)
    sources = search_documents(question, root=root)
    above = [item for item in validation.screening if item.classification == "above_criterion"]
    source_lines = [f"- [{item['title']}]({item['source_uri']})" for item in sources["matches"]]
    finding_lines = [
        f"- {item.sample_id}: {item.analyte} = {item.result:g} {item.unit} "
        f"(screening criterion {item.criterion:g} {item.criterion_unit})"
        for item in above
    ]
    report = "\n".join(
        [
            f"# Draft screening summary — {project_id}",
            "",
            "## Scope",
            f"Reviewed `{lab_file}` against demonstration screening criteria for: {question}",
            "",
            "## Data quality",
            f"Parsed {validation.row_count} result rows. Validation status: {'passed' if validation.valid else 'failed'}.",
            *(f"- {error}" for error in validation.errors),
            "",
            "## Screening observations",
            *(finding_lines or ["- No results were above the available demonstration criteria."]),
            "",
            "## Sources",
            *(source_lines or ["- No authorized source matched the question; practitioner review is required."]),
            "",
            f"> {DISCLAIMER}",
        ]
    )
    return {
        "status": "draft_requires_human_review",
        "project_id": project_id,
        "report_markdown": report,
        "validation": validation.model_dump(mode="json"),
        "sources": sources["matches"],
        "can_publish": False,
    }
