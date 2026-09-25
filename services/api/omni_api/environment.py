# ruff: noqa: B008
import asyncio
import hashlib
from contextlib import suppress
from datetime import UTC, datetime
from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.shared import RGBColor
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from services.api.omni_api.audit import add_audit_event
from services.api.omni_api.auth import TenantContext, get_tenant_context, require_roles
from services.api.omni_api.database import get_session
from services.api.omni_api.environment_schemas import (
    EnvironmentalProjectCreate,
    EnvironmentalProjectRead,
    EnvironmentalReportCreate,
    EnvironmentalReportRead,
    EnvironmentalReportReview,
    EnvironmentalSearchResult,
    EnvironmentalWorkbookRead,
)
from services.api.omni_api.models import (
    EnvironmentalProject,
    EnvironmentalReport,
    EnvironmentalWorkbook,
)
from services.api.omni_api.storage import AttachmentStorage
from services.environment_mcp.domain import (
    DISCLAIMER,
    draft_screening_report_from_validation,
    search_documents,
    validate_lab_workbook_bytes,
)

router = APIRouter(prefix="/environment", tags=["environment"])


async def _project(session: AsyncSession, tenant_id: str, project_id: str) -> EnvironmentalProject:
    record = await session.scalar(
        select(EnvironmentalProject).where(
            EnvironmentalProject.id == project_id,
            EnvironmentalProject.tenant_id == tenant_id,
        )
    )
    if record is None:
        raise HTTPException(status_code=404, detail="Environmental project not found")
    return record


async def _workbook(session: AsyncSession, tenant_id: str, workbook_id: str) -> EnvironmentalWorkbook:
    record = await session.scalar(
        select(EnvironmentalWorkbook).where(
            EnvironmentalWorkbook.id == workbook_id,
            EnvironmentalWorkbook.tenant_id == tenant_id,
        )
    )
    if record is None:
        raise HTTPException(status_code=404, detail="Environmental workbook not found")
    return record


async def _report(session: AsyncSession, tenant_id: str, report_id: str) -> EnvironmentalReport:
    record = await session.scalar(
        select(EnvironmentalReport).where(
            EnvironmentalReport.id == report_id,
            EnvironmentalReport.tenant_id == tenant_id,
        )
    )
    if record is None:
        raise HTTPException(status_code=404, detail="Environmental report not found")
    return record


@router.get("/projects", response_model=list[EnvironmentalProjectRead])
async def list_projects(
    context: TenantContext = Depends(get_tenant_context),
    session: AsyncSession = Depends(get_session),
):
    records = await session.scalars(
        select(EnvironmentalProject)
        .where(EnvironmentalProject.tenant_id == context.tenant_id)
        .order_by(EnvironmentalProject.created_at.desc())
    )
    return list(records)


@router.post("/projects", response_model=EnvironmentalProjectRead, status_code=201)
async def create_project(
    payload: EnvironmentalProjectCreate,
    request: Request,
    context: TenantContext = Depends(require_roles("owner", "admin", "reviewer")),
    session: AsyncSession = Depends(get_session),
):
    record = EnvironmentalProject(
        tenant_id=context.tenant_id,
        created_by_user_id=context.user.id,
        **payload.model_dump(),
    )
    session.add(record)
    try:
        await session.flush()
    except IntegrityError:
        await session.rollback()
        raise HTTPException(status_code=409, detail="Project code already exists for this tenant") from None
    add_audit_event(
        session,
        tenant_id=context.tenant_id,
        actor_user_id=context.user.id,
        action="environment.project.created",
        resource_type="environmental_project",
        resource_id=record.id,
        correlation_id=request.state.correlation_id,
        payload={"code": record.code},
    )
    await session.commit()
    await session.refresh(record)
    return record


@router.get("/projects/{project_id}/workbooks", response_model=list[EnvironmentalWorkbookRead])
async def list_workbooks(
    project_id: str,
    context: TenantContext = Depends(get_tenant_context),
    session: AsyncSession = Depends(get_session),
):
    await _project(session, context.tenant_id, project_id)
    records = await session.scalars(
        select(EnvironmentalWorkbook)
        .where(
            EnvironmentalWorkbook.tenant_id == context.tenant_id,
            EnvironmentalWorkbook.project_id == project_id,
        )
        .order_by(EnvironmentalWorkbook.created_at.desc())
    )
    return list(records)


@router.post("/projects/{project_id}/workbooks", response_model=EnvironmentalWorkbookRead, status_code=201)
async def upload_workbook(
    project_id: str,
    request: Request,
    file: UploadFile = File(...),
    context: TenantContext = Depends(require_roles("owner", "admin", "reviewer")),
    session: AsyncSession = Depends(get_session),
):
    project = await _project(session, context.tenant_id, project_id)
    filename = (file.filename or "laboratory-results.xlsx").split("/")[-1].split("\\")[-1]
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="Laboratory upload must be an .xlsx workbook")
    content = await file.read(20 * 1024 * 1024 + 1)
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Laboratory workbook exceeds the 20 MB safety limit")
    if not content:
        raise HTTPException(status_code=422, detail="Laboratory workbook is empty")

    validation = await asyncio.to_thread(validate_lab_workbook_bytes, project.code, filename, content)
    record = EnvironmentalWorkbook(
        tenant_id=context.tenant_id,
        project_id=project.id,
        uploaded_by_user_id=context.user.id,
        filename=filename,
        content_type=file.content_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=len(content),
        object_key="pending",
        content_hash=hashlib.sha256(content).hexdigest(),
        status="validated" if validation.valid else "invalid",
        validation=validation.model_dump(mode="json"),
    )
    session.add(record)
    await session.flush()
    record.object_key = f"environment/{context.tenant_id}/{project.id}/{record.id}.xlsx"
    storage = AttachmentStorage(request.app.state.settings)
    try:
        await asyncio.to_thread(storage.put, record.object_key, content, record.content_type)
        add_audit_event(
            session,
            tenant_id=context.tenant_id,
            actor_user_id=context.user.id,
            action="environment.workbook.uploaded",
            resource_type="environmental_workbook",
            resource_id=record.id,
            correlation_id=request.state.correlation_id,
            payload={"filename": filename, "status": record.status, "content_hash": record.content_hash},
        )
        await session.commit()
    except Exception:
        await session.rollback()
        with suppress(Exception):
            await asyncio.to_thread(storage.delete, record.object_key)
        raise
    await session.refresh(record)
    return record


@router.get("/documents/search", response_model=EnvironmentalSearchResult)
async def document_search(
    q: str = Query(min_length=3, max_length=500),
    _: TenantContext = Depends(get_tenant_context),
):
    return await asyncio.to_thread(search_documents, q)


@router.get("/projects/{project_id}/reports", response_model=list[EnvironmentalReportRead])
async def list_reports(
    project_id: str,
    context: TenantContext = Depends(get_tenant_context),
    session: AsyncSession = Depends(get_session),
):
    await _project(session, context.tenant_id, project_id)
    records = await session.scalars(
        select(EnvironmentalReport)
        .where(
            EnvironmentalReport.tenant_id == context.tenant_id,
            EnvironmentalReport.project_id == project_id,
        )
        .order_by(EnvironmentalReport.created_at.desc())
    )
    return list(records)


@router.post("/projects/{project_id}/reports", response_model=EnvironmentalReportRead, status_code=201)
async def create_report(
    project_id: str,
    payload: EnvironmentalReportCreate,
    request: Request,
    context: TenantContext = Depends(require_roles("owner", "admin", "reviewer")),
    session: AsyncSession = Depends(get_session),
):
    project = await _project(session, context.tenant_id, project_id)
    workbook = await _workbook(session, context.tenant_id, payload.workbook_id)
    if workbook.project_id != project.id:
        raise HTTPException(status_code=422, detail="Workbook does not belong to this project")
    if workbook.status != "validated":
        raise HTTPException(status_code=422, detail="Workbook validation must pass before drafting a report")
    draft = await asyncio.to_thread(
        draft_screening_report_from_validation,
        project.code,
        workbook.filename,
        payload.question,
        workbook.validation,
    )
    record = EnvironmentalReport(
        tenant_id=context.tenant_id,
        project_id=project.id,
        workbook_id=workbook.id,
        created_by_user_id=context.user.id,
        question=payload.question,
        report_markdown=draft["report_markdown"],
        sources=draft["sources"],
        status="draft",
    )
    session.add(record)
    await session.flush()
    add_audit_event(
        session,
        tenant_id=context.tenant_id,
        actor_user_id=context.user.id,
        action="environment.report.drafted",
        resource_type="environmental_report",
        resource_id=record.id,
        correlation_id=request.state.correlation_id,
        payload={"workbook_id": workbook.id, "source_count": len(record.sources)},
    )
    await session.commit()
    await session.refresh(record)
    return record


@router.post("/reports/{report_id}/review", response_model=EnvironmentalReportRead)
async def review_report(
    report_id: str,
    payload: EnvironmentalReportReview,
    request: Request,
    context: TenantContext = Depends(require_roles("owner", "admin", "reviewer")),
    session: AsyncSession = Depends(get_session),
):
    report = await _report(session, context.tenant_id, report_id)
    report.status = "approved" if payload.decision == "approve" else "rejected"
    report.review_reason = payload.reason
    report.reviewed_by_user_id = context.user.id
    report.reviewed_at = datetime.now(UTC)
    if payload.corrected_markdown:
        report.report_markdown = payload.corrected_markdown
        report.version += 1
    add_audit_event(
        session,
        tenant_id=context.tenant_id,
        actor_user_id=context.user.id,
        action=f"environment.report.{report.status}",
        resource_type="environmental_report",
        resource_id=report.id,
        correlation_id=request.state.correlation_id,
        payload={"reason": payload.reason, "version": report.version},
    )
    await session.commit()
    await session.refresh(report)
    return report


def _docx_export(report: EnvironmentalReport, project: EnvironmentalProject) -> bytes:
    document = Document()
    if report.status != "approved":
        warning = document.add_heading("DRAFT — NOT APPROVED FOR CLIENT OR REGULATORY USE", level=1)
        for run in warning.runs:
            run.font.color.rgb = RGBColor(192, 35, 24)
    document.add_heading(project.name, level=0)
    document.add_paragraph(f"Project: {project.code} · Jurisdiction: {project.jurisdiction}")
    document.add_paragraph(f"Report status: {report.status.upper()} · Version: {report.version}")
    for line in report.report_markdown.splitlines():
        text = line.strip()
        if text.startswith("# "):
            document.add_heading(text[2:], level=1)
        elif text.startswith("## "):
            document.add_heading(text[3:], level=2)
        elif text.startswith("- "):
            document.add_paragraph(text[2:], style="List Bullet")
        elif text:
            document.add_paragraph(text.lstrip("> "))
    document.add_paragraph(DISCLAIMER)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_export(report: EnvironmentalReport, project: EnvironmentalProject, workbook: EnvironmentalWorkbook) -> bytes:
    def safe_cell(value):
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    output_book = Workbook()
    summary = output_book.active
    summary.title = "Report summary"
    summary.append(
        ["DRAFT — NOT APPROVED FOR CLIENT OR REGULATORY USE"] if report.status != "approved" else ["APPROVED"]
    )
    summary.append(["Project", project.name])
    summary.append(["Project code", project.code])
    summary.append(["Jurisdiction", project.jurisdiction])
    summary.append(["Status", report.status])
    summary.append(["Question", report.question])
    summary.append(["Disclaimer", DISCLAIMER])
    screening = output_book.create_sheet("Screening results")
    screening.append(
        ["Sample ID", "Analyte", "Result", "Unit", "Criterion", "Criterion unit", "Classification", "Source"]
    )
    for item in workbook.validation.get("screening", []):
        screening.append(
            [
                safe_cell(item.get("sample_id")),
                safe_cell(item.get("analyte")),
                item.get("result"),
                safe_cell(item.get("unit")),
                item.get("criterion"),
                safe_cell(item.get("criterion_unit")),
                safe_cell(item.get("classification")),
                safe_cell(item.get("source_uri")),
            ]
        )
    output = BytesIO()
    output_book.save(output)
    return output.getvalue()


@router.get("/reports/{report_id}/export.{format}")
async def export_report(
    report_id: str,
    format: str,
    request: Request,
    context: TenantContext = Depends(get_tenant_context),
    session: AsyncSession = Depends(get_session),
):
    if format not in {"docx", "xlsx"}:
        raise HTTPException(status_code=404, detail="Export format not found")
    report = await _report(session, context.tenant_id, report_id)
    project = await _project(session, context.tenant_id, report.project_id)
    workbook = await _workbook(session, context.tenant_id, report.workbook_id)
    if format == "docx":
        content = await asyncio.to_thread(_docx_export, report, project)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        content = await asyncio.to_thread(_xlsx_export, report, project, workbook)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    add_audit_event(
        session,
        tenant_id=context.tenant_id,
        actor_user_id=context.user.id,
        action="environment.report.exported",
        resource_type="environmental_report",
        resource_id=report.id,
        correlation_id=request.state.correlation_id,
        payload={"format": format, "status": report.status},
    )
    await session.commit()
    filename = quote(f"{project.code}-{report.id[:8]}-{report.status}.{format}")
    return StreamingResponse(
        BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
